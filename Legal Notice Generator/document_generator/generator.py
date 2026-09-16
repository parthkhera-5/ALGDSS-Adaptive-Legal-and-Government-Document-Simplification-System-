import re
import uuid
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from config.config import (TEMPLATE_DIR,FILLED_DIR,BLANK_DIR)
# ============================================================
# TEMPLATE LOADING
# ============================================================
def load_docx_template(template_name):
    template_path = (TEMPLATE_DIR / template_name)
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")
    return Document(str(template_path))
# ============================================================
# REPLACE PLACEHOLDERS WITH BLANK LINES
# ============================================================
#
# Used ONLY for blank documents.
#
# Example:
#
# {{name}}
#
# becomes:
#
# ________________________________
#
# ============================================================
# def replace_placeholders_with_lines(document):
#     pattern = re.compile(r"\{\{([^{}]+)\}\}")
#     def process_paragraph(paragraph):
#         # ----------------------------------------------------
#         # Combine all runs
#         # ----------------------------------------------------
#         full_text = "".join(run.text or "" for run in paragraph.runs)
#         # ----------------------------------------------------
#         # No placeholder
#         # ----------------------------------------------------
#         if not pattern.search(full_text):
#             return
#         # ----------------------------------------------------
#         # Replace every {{field}} with 32 underscores
#         # ----------------------------------------------------
#         new_text = pattern.sub(lambda match: "_" * 32,full_text)
#         # ----------------------------------------------------
#         # Put complete text into first run
#         # ----------------------------------------------------
#         if paragraph.runs:
#             paragraph.runs[0].text = new_text
#             # Clear remaining runs
#             for run in paragraph.runs[1:]:
#                 run.text = ""
#     # --------------------------------------------------------
#     # Normal paragraphs
#     # --------------------------------------------------------
#     for paragraph in document.paragraphs:
#         process_paragraph(paragraph)
#     # --------------------------------------------------------
#     # Tables
#     # --------------------------------------------------------\
#     for table in document.tables:
#         for row in table.rows:
#             for cell in row.cells:
#                 for paragraph in cell.paragraphs:
#                     process_paragraph(paragraph)
#     return document



def replace_placeholders_with_lines(document):
    """
    Replace both {{field}} and {field} with blank lines.
    """
    pattern = re.compile(r"\{\{([^{}]+)\}\}|\{([^{}]+)\}")
    def process_paragraph(paragraph):
        full_text = "".join(run.text or "" for run in paragraph.runs)
        if not pattern.search(full_text):
            return
        new_text = pattern.sub("_" * 32, full_text)
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
    for paragraph in document.paragraphs:
        process_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)
    return document
# ============================================================
# REPLACE PLACEHOLDERS WITH USER VALUES
# ============================================================
#
# Used ONLY for filled documents.
#
# Example:
#
# {{name}}
#
# becomes:
#
# Parth Khera
#
# ============================================================
# def replace_placeholders_with_values(document,data):
#     pattern = re.compile(r"\{\{([^{}]+)\}\}")
#     def process_paragraph(paragraph):
#         # ----------------------------------------------------
#         # Combine all runs
#         # ----------------------------------------------------
#         full_text = "".join(run.text or "" for run in paragraph.runs)
#         # ----------------------------------------------------
#         # No placeholder
#         # ----------------------------------------------------
#         if not pattern.search(full_text):
#             return
#         # ----------------------------------------------------
#         # Replace placeholders
#         # ----------------------------------------------------
#         def replace_match(match):
#             field_name = (match.group(1).strip())
#             value = data.get(field_name)
#             # ------------------------------------------------
#             # If value exists
#             # ------------------------------------------------
#             if value is not None:
#                 return str(value)
#             # ------------------------------------------------
#             # If value does not exist, keep placeholder
#             # ------------------------------------------------
#             return match.group(0)
#         new_text = pattern.sub(replace_match,full_text)
#         # ----------------------------------------------------
#         # Put complete text into first run
#         # ----------------------------------------------------
#         if paragraph.runs:
#             paragraph.runs[0].text = new_text
#             # Clear remaining runs
#             for run in paragraph.runs[1:]:
#                 run.text = ""
#     # --------------------------------------------------------
#     # Normal paragraphs
#     # --------------------------------------------------------
#     for paragraph in document.paragraphs:
#         process_paragraph(paragraph)
#     # --------------------------------------------------------
#     # Tables
#     # --------------------------------------------------------
#     for table in document.tables:
#         for row in table.rows:
#             for cell in row.cells:
#                 for paragraph in cell.paragraphs:
#                     process_paragraph(paragraph)
#     return document


def replace_placeholders_with_values(document, data):
    """
    Replace both {{field}} and {field} placeholders with user values.
    """
    pattern = re.compile(r"\{\{([^{}]+)\}\}|\{([^{}]+)\}")
    def process_paragraph(paragraph):
        full_text = "".join(run.text or "" for run in paragraph.runs)
        if not pattern.search(full_text):
            return
        def replace_match(match):
            field_name = (match.group(1) or match.group(2)).strip()
            value = data.get(field_name)
            if value is not None:
                return str(value)
            return match.group(0)
        new_text = pattern.sub(replace_match, full_text)
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
    for paragraph in document.paragraphs:
        process_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)
    return document
# ============================================================
# GENERATE FILLED DOCUMENT FROM TEMPLATE
# ============================================================
def generate_from_template(template_name,data,output_name=None):
    # --------------------------------------------------------
    # Load template
    # --------------------------------------------------------
    document = load_docx_template(template_name)
    # --------------------------------------------------------
    # Replace placeholders with actual values
    # --------------------------------------------------------
    document = replace_placeholders_with_values(document,data)
    # --------------------------------------------------------
    # Generate output filename
    # --------------------------------------------------------
    if output_name is None:
        output_name = (f"legal_document_" f"{uuid.uuid4().hex[:8]}.docx")
    # --------------------------------------------------------
    # Output path
    # --------------------------------------------------------
    output_path = (FILLED_DIR / output_name)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    # --------------------------------------------------------
    # Save filled document
    # --------------------------------------------------------
    document.save(str(output_path))
    print(f"Filled document generated: {output_path}")
    return output_path
# ============================================================
# CREATE COMPLETELY BLANK DOCUMENT
# ============================================================
def create_blank_document(title="Legal Document",output_name=None):
    document = Document()
    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------
    title_paragraph = (document.add_paragraph())
    title_paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER)
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    # --------------------------------------------------------
    # Empty line
    # --------------------------------------------------------
    document.add_paragraph("")
    # --------------------------------------------------------
    # Filename
    # --------------------------------------------------------
    if output_name is None:
        output_name = (f"blank_" f"{uuid.uuid4().hex[:8]}.docx")
    # --------------------------------------------------------
    # Output path
    # --------------------------------------------------------
    output_path = (BLANK_DIR / output_name)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------
    document.save(str(output_path))
    return output_path
# ============================================================
# CREATE DOCUMENT FROM AI GENERATED TEXT
# ============================================================
def create_document_from_text(text,title="Legal Document",output_name=None):
    document = Document()
    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------
    title_paragraph = (document.add_paragraph())
    title_paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER)
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    # --------------------------------------------------------
    # Empty line
    # --------------------------------------------------------
    document.add_paragraph("")
    # --------------------------------------------------------
    # Add generated text
    # --------------------------------------------------------
    for paragraph_text in text.split("\n"):
        if paragraph_text.strip():
            document.add_paragraph(paragraph_text.strip())
    # --------------------------------------------------------
    # Filename
    # --------------------------------------------------------
    if output_name is None:
        output_name = (f"generated_" f"{uuid.uuid4().hex[:8]}.docx")
    # --------------------------------------------------------
    # Output path
    # --------------------------------------------------------
    output_path = (FILLED_DIR / output_name)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------
    document.save(str(output_path))
    return output_path