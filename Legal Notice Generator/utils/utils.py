import os
import re
import uuid
import pymupdf
from docx import Document

# ============================================================
# ALLOWED FILE
# ============================================================
def allowed_file(filename,allowed_extensions):
    return ("." in filename and filename.rsplit(".", 1)[1].lower() in allowed_extensions)

# ============================================================
# GENERATE ID
# ============================================================
def generate_id():
    return uuid.uuid4().hex
# ============================================================
# CLEAN TEXT
# ============================================================
def clean_text(text):
    if not text:
        return ""
    text = re.sub(r"\s+"," ",text)
    return text.strip()
# ============================================================
# EXTRACT PDF TEXT
# ============================================================
def extract_pdf_text(file_path):
    text_parts = []
    document = pymupdf.open(file_path)
    try:
        for page in document:
            text_parts.append(page.get_text())
    finally:
        document.close()
    return clean_text("\n".join(text_parts))
# ============================================================
# EXTRACT DOCX TEXT
# ============================================================
def extract_docx_text(file_path):
    document = Document(file_path)
    text_parts = []
    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)
    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------
    for table in document.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                text_parts.append(" | ".join(row_text))
    return clean_text("\n".join(text_parts))
# ============================================================
# EXTRACT TEXT
# ============================================================
def extract_text(file_path):
    extension = (os.path.splitext(file_path)[1].lower())
    if extension == ".pdf":
        return extract_pdf_text(file_path)
    if extension == ".docx":
        return extract_docx_text(file_path)
    if extension == ".txt":
        with open(file_path,"r",encoding="utf-8") as file:
            return clean_text(file.read())
    raise ValueError(f"Unsupported file type: {extension}")